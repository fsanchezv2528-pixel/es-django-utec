"""Suite de pruebas W12 — PDF de factura y exportación Excel.

Funciona en Windows sin Docker (xhtml2pdf) y dentro de Docker (WeasyPrint).

Ejecutar con:
    python manage.py test tests.test_w12_documentos --verbosity=2

Resultado esperado: Ran 8 tests ... OK
"""
import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from clientes.models  import Cliente
from productos.models import Categoria, Producto
from ventas.models    import DetalleVenta, Venta


class DocumentosTest(TestCase):

    def setUp(self):
        self.user = User.objects.create_user('pdfuser', password='pass')
        cat  = Categoria.objects.create(nombre='Cat PDF')
        prod = Producto.objects.create(
            nombre='Producto PDF', precio=Decimal('500.00'), stock=5, categoria=cat
        )
        cli = Cliente.objects.create(nombre='Cliente PDF', correo='pdf@test.com')
        self.venta = Venta.objects.create(cliente=cli)
        DetalleVenta.objects.create(
            venta=self.venta, producto=prod, cantidad=2, precio_unitario=Decimal('500.00')
        )

    def _pdf(self):
        self.client.force_login(self.user)
        return self.client.get(reverse('ventas:pdf', args=[self.venta.pk]))

    def test_pdf_sin_auth_redirige_a_login(self):
        r = self.client.get(reverse('ventas:pdf', args=[self.venta.pk]))
        self.assertEqual(r.status_code, 302)
        self.assertIn('/accounts/login/', r['Location'])

    def test_pdf_con_auth_devuelve_200(self):
        self.assertEqual(self._pdf().status_code, 200)

    def test_pdf_content_type_application_pdf(self):
        self.assertIn('application/pdf', self._pdf().get('Content-Type', ''))

    def test_pdf_empieza_con_header_pdf(self):
        self.assertEqual(self._pdf().content[:4], b'%PDF')

    def test_pdf_content_disposition_attachment(self):
        disposition = self._pdf().get('Content-Disposition', '')
        self.assertIn('attachment', disposition)
        self.assertIn('.pdf', disposition)

    def test_pdf_venta_inexistente_devuelve_404(self):
        self.client.force_login(self.user)
        r = self.client.get(reverse('ventas:pdf', args=[9999]))
        self.assertEqual(r.status_code, 404)

    def test_pdf_es_mayor_que_1kb(self):
        self.assertGreater(len(self._pdf().content), 1024)

    def test_excel_exportacion_funciona(self):
        self.client.force_login(self.user)
        r = self.client.get(reverse('productos:exportar_excel'))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.get('Content-Type', ''))
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws['B1'].value, 'Nombre')
        self.assertEqual(ws['B2'].value, 'Producto PDF')
